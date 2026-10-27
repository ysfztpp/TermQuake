"""The sync step without the network: xlsx bytes -> CSV snapshots."""

import io

import pytest
from openpyxl import Workbook

from termquake.fetch import resolve_tab, workbook_to_csvs


def make_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "All Terms and Expressions Terim"  # how xlsx truncates the real tab name
    ws.append(["Legend", None, None])
    ws.merge_cells("A1:A2")
    ws.append([None, None, None])
    ws.append(["Nr.", "English", "Türkçe (Turkish)"])
    ws.append([1.0, "earthquake", "deprem"])
    ws.append([2, "rubble", None, None, None])  # trailing empty cells
    ws.append([None, None, None])  # trailing empty row
    other = wb.create_sheet("Terms Needing Definition Selection")
    other.append(["", "English", True, 0.8776])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_workbook_to_csvs():
    tabs = {
        "All Terms and Expressions /Terimler ve İfadeler": "terms",
        "Terms Needing Definition Selection": "definitions",
    }
    out = workbook_to_csvs(make_xlsx(), tabs)
    assert (
        out["terms"]
        == "Legend,,\n,,\nNr.,English,Türkçe (Turkish)\n1,earthquake,deprem\n2,rubble,\n"
    )
    assert out["definitions"] == ",English,TRUE,0.8776\n"


@pytest.mark.parametrize(
    ("wanted", "available", "expected"),
    [
        ("Term submission", ["Term submission", "TEST"], "Term submission"),
        (
            "All Terms and Expressions /Terimler ve İfadeler",
            ["All Terms and Expressions Terim", "Source Term Selection"],
            "All Terms and Expressions Terim",
        ),
        (
            "All Terms and Expressions /Terimler ve İfadeler",
            ["All Terms and Expressions -Terim"[:31]],
            None,
        ),
    ],
)
def test_resolve_tab(wanted, available, expected):
    if expected is None:
        with pytest.raises(RuntimeError):
            resolve_tab(wanted, available)
    else:
        assert resolve_tab(wanted, available) == expected


def test_renamed_tab_fails_with_clear_message():
    with pytest.raises(RuntimeError, match="termquake.toml"):
        workbook_to_csvs(make_xlsx(), {"Old tab name": "terms"})
