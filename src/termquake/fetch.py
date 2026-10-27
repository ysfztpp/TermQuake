"""Download a Google Sheet (read-only) and save selected tabs as CSV.

Used once to snapshot the legacy sheet into data/legacy/ before the
migration. Normal work never calls this: the SQL database is the source of
truth. Kept for re-importing from a *copy* of the sheet if ever needed
(`termquake fetch-sheet --sheet-id <ID of the copy>`).

The xlsx export of a sheet shared as "anyone with the link can view" needs no
API key and returns every tab in one request. Downloading never modifies the sheet.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

USER_AGENT = "TermQuake-sync/1.0 (+https://github.com/ysfztpp/termquake)"


def download_xlsx(url: str, timeout: int = 60) -> bytes:
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=timeout) as response:  # noqa: S310 (fixed https URL)
        data = response.read()
    if not data.startswith(b"PK"):
        raise RuntimeError(
            "Download is not an xlsx file. Is the sheet still shared as "
            "'anyone with the link can view'?"
        )
    return data


def _cell_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime | dt.date):
        return value.isoformat()
    return str(value)


_XLSX_FORBIDDEN = re.compile(r"[\[\]:*?/\\\s]")


def _tab_key(name: str) -> str:
    return _XLSX_FORBIDDEN.sub("", name).casefold()


def resolve_tab(wanted: str, available: list[str]) -> str:
    """Find a Google Sheets tab inside the xlsx export.

    xlsx tab names cannot contain / \\ ? * [ ] : and are cut at 31 characters,
    so Google renames e.g. "All Terms and Expressions /Terimler ve İfadeler"
    on export. Match exactly first, then ignoring forbidden characters and
    allowing truncation.
    """
    if wanted in available:
        return wanted
    key = _tab_key(wanted)
    hits = [name for name in available if (k := _tab_key(name)) and key.startswith(k)]
    if len(hits) == 1:
        return hits[0]
    raise RuntimeError(
        f"Tab {wanted!r} not found (candidates: {hits or 'none'}). Available tabs: {available}. "
        "Update [sheet.tabs] in termquake.toml if a tab was renamed."
    )


def workbook_to_csvs(xlsx: bytes, tabs: dict[str, str]) -> dict[str, str]:
    """Return {file_stem: csv_text} for the configured tabs."""
    wb = load_workbook(io.BytesIO(xlsx), read_only=True, data_only=True)
    out: dict[str, str] = {}
    try:
        for tab, stem in tabs.items():
            sheet = wb[resolve_tab(tab, wb.sheetnames)]
            rows = [[_cell_to_text(v) for v in row] for row in sheet.iter_rows(values_only=True)]
            out[stem] = _rows_to_csv(rows)
    finally:
        wb.close()
    return out


def _rows_to_csv(rows: list[list[str]]) -> str:
    # Trim trailing empty rows and columns so the snapshot only changes when content does.
    while rows and not any(cell.strip() for cell in rows[-1]):
        rows.pop()
    width = 0
    for row in rows:
        for i in range(len(row) - 1, -1, -1):
            if row[i].strip():
                width = max(width, i + 1)
                break
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    for row in rows:
        writer.writerow((row + [""] * width)[:width])
    return buf.getvalue()


def fetch_sheet(sheet_id: str, tabs: dict[str, str], out_dir: Path) -> list[str]:
    """Download a sheet and write <out_dir>/<stem>.csv per tab. Returns files that changed."""
    out_dir.mkdir(parents=True, exist_ok=True)
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    csvs = workbook_to_csvs(download_xlsx(url), tabs)
    changed: list[str] = []
    for stem, text in csvs.items():
        path = out_dir / f"{stem}.csv"
        old = path.read_text(encoding="utf-8") if path.exists() else None
        if old != text:
            path.write_text(text, encoding="utf-8")
            changed.append(path.name)
    return changed
